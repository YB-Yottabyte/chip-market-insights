import pandas as pd
import pytest
from openpyxl import Workbook

from src.load_industry import load_directory
from src.export_powerbi import build_onsemi_industry_comparison


def test_tidy_industry_file_is_standardized(tmp_path):
    file = tmp_path / "billings.csv"
    file.write_text("Date,Region,Monthly Sales\n2025-01-01,APAC,100\n")
    frame = load_directory(tmp_path)
    assert frame.iloc[0].region == "Asia Pacific"
    assert frame.iloc[0].date == "2025-01-31"
    assert frame.iloc[0].monthly_sales == 100


def test_duplicate_month_rejected(tmp_path):
    (tmp_path / "a.csv").write_text("Date,Region,Monthly Sales\n2025-01-01,Japan,10\n")
    (tmp_path / "b.csv").write_text("Date,Region,Monthly Sales\n2025-01-31,Japan,11\n")
    with pytest.raises(ValueError):
        load_directory(tmp_path)


def test_excel_input_uses_same_loader(tmp_path):
    pd.DataFrame(
        {
            "Date": ["2025-02-01"],
            "Region": ["Europe"],
            "Three-Month Moving Average": [42.5],
        }
    ).to_excel(tmp_path / "billings.xlsx", index=False)
    result = load_directory(tmp_path)
    assert result.iloc[0].three_month_average == 42.5
    assert pd.isna(result.iloc[0].monthly_sales)


def test_native_wsts_workbook_converts_thousands_usd(tmp_path):
    book = Workbook()
    monthly = book.active
    monthly.title = "Monthly Data"
    average = book.create_sheet("3MMA")
    for sheet in (monthly, average):
        sheet.cell(3, 1, "All numbers are in 1000 US$.")
        for month in range(1, 13):
            sheet.cell(4, month + 1, pd.Timestamp(2025, month, 1).strftime("%B"))
        sheet.cell(5, 1, 2025)
        sheet.cell(6, 1, "Worldwide")
    monthly.cell(6, 2, 42)
    average.cell(6, 2, 40)
    average.cell(6, 3, 0)  # Future placeholder is not a February record.
    path = tmp_path / "wsts.xlsx"
    book.save(path)
    result = load_directory(tmp_path)
    assert len(result) == 1
    assert result.iloc[0].region == "World"
    assert result.iloc[0].monthly_sales == 42000
    assert result.iloc[0].three_month_average == 40000


def test_comparison_requires_complete_market_windows():
    company = pd.DataFrame(
        {
            "ticker": ["ON", "ON"],
            "period_end": ["2024-03-31", "2025-03-31"],
            "fiscal_year": [2024, 2025],
            "fiscal_quarter": [1, 1],
            "revenue": [100, 120],
            "gross_profit": [40, 50],
            "operating_income": [20, 25],
            "net_income": [10, 12],
            "rd_expense": [5, 6],
        }
    )
    markets = pd.DataFrame(
        {
            "region": ["World"] * 6,
            "date": [
                "2024-01-31",
                "2024-02-29",
                "2024-03-31",
                "2025-01-31",
                "2025-02-28",
                "2025-03-31",
            ],
            "monthly_sales": [10, 10, 10, 12, 12, 12],
            "source_file": ["official.csv"] * 6,
        }
    )
    result = build_onsemi_industry_comparison(company, markets)
    assert len(result) == 1
    assert abs(result.iloc[0].growth_gap) < 1e-10
    incomplete = markets.drop(index=1)
    assert build_onsemi_industry_comparison(company, incomplete).empty
